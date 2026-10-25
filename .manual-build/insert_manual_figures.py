from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
TARGET = ROOT / "User-Manual.docx"
SOURCE = ROOT / "Test-System-Screenshots" / "Screenshots"
IMAGE_DIR = ROOT / "User-Manual-Images"

FIGURES = [
    ("tc-stk-03_stock-documents-and-reports-rows-open-and-back.png", "Result: You work inside the selected store.", "Tablet Stock screen with the side rail and grouped work areas visible."),
    ("tc-str-04_store-creation-levels_step-2_80-character-name.png", "Result: The store appears on Home.", "Level 1 selected in the second step of store creation."),
    ("tc-itm-20_inventory-ocr-scan_level-3-expanded-item-form.png", "Result: The item appears in the Items list.", "New Item form with the main item details and selling-price fields."),
    ("tc-itp-11_item-prices-current-catalog-rate-list.png", "Result: The list shows the item, rate, price list and dates.", "Saved item prices listed by item, price list, unit, rate and validity."),
    ("tc-imi-01_import-items_page-template-columns.png", "Result: Items appear in the Items list after import.", "Import Items field guide shown before a CSV file is selected."),
    ("tc-pos-03_pos-transactions_session-opened-with-500.png", "Result: Sell shows the open session, register and date,", "Open session card with Main Register and opening cash shown."),
    ("tc-pos-10_pos-transactions_register-row-adds-to-cart.png", "Result: Selected items appear as cart lines.", "Tablet Register with the item list beside the current Cart."),
    ("tc-pos-12_pos-transactions_overpayment-change-90.png", "5. Check Change, the money to give back to the customer.", "Payment review with cash received and change before completing the order."),
    ("tc-pos-14_pos-transactions_double-complete-one-sale.png", "Result: The sale is completed and its applicable stock", "Completed order receipt with totals and ledger links."),
    ("tc-pos-21_pos-transactions_orders-list-submitted.png", "Result: You can review the saved sale and reopen its receipt.", "Orders list showing submitted, returned and cancelled records."),
    ("tc-pos-23_pos-transactions_partial-refund-receipt.png", "Result: The return receipt opens.", "Refund receipt showing the returned item and cash refund."),
    ("tc-pos-25_pos-transactions_close-session-after-sale-return-count-1250.png", "Result: The session closes and its closing record is available.", "POS Closing Entry with expected cash, the count and the difference."),
    ("tc-dsh-01_dashboard_open-and-scroll.png", "Result: You see a summary of the records you are permitted to read.", "Tablet Dashboard with summary cards, sales chart and top-selling items."),
    ("tc-emp-02_three-default-roles.png", "Result: The role appears with a coloured dot", "Built-in roles with their employee and permission counts."),
    ("tc-ftr-02_hide-show-all-setup-pages-restored.png", "Result: The page confirms Saved.", "Features screen with store screens on the left and visibility switches on the right."),
    ("tc-gde-01_store-guide-role-picker-opens.png", "Result: You practise with sample items, sales and sessions", "Store Guide role picker for choosing a role to practise."),
    ("tc-cmp-09_company-accounts-and-stock-settings.png", "Result: The store's identity and fallback accounts are updated.", "Store settings with fallback accounts and perpetual-inventory controls."),
    ("tc-mop-01_modes-of-payment_initial-list.png", "Result: The mode can be selected in POS Profiles", "Modes of Payment list showing method, default account and type."),
    ("tc-prf-08_print-formats_pos-receipt-initial-default.png", "Result: Receipts use the profile's format or the default.", "Print Formats list with POS Receipt selected as the default."),
]


def matching_paragraph(document, prefix):
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting with {prefix!r}; found {len(matches)}")
    return matches[0]


def add_figure_after(anchor, image_path, number, caption):
    image_p = Paragraph(OxmlElement("w:p"), anchor._parent)
    anchor._p.addnext(image_p._p)
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_before = Pt(5)
    image_p.paragraph_format.space_after = Pt(0)
    image_p.paragraph_format.keep_with_next = True
    image_p.paragraph_format.keep_together = True
    run = image_p.add_run()
    shape = run.add_picture(str(image_path), width=Inches(6.55))
    shape._inline.docPr.set("title", f"Figure {number}. {caption}")
    shape._inline.docPr.set("descr", caption)

    caption_p = Paragraph(OxmlElement("w:p"), image_p._parent)
    image_p._p.addnext(caption_p._p)
    caption_p.style = "Caption"
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(2)
    caption_p.paragraph_format.space_after = Pt(9)
    caption_p.paragraph_format.keep_together = True
    caption_run = caption_p.add_run(f"Figure {number}. {caption}")
    caption_run.font.name = "Calibri"
    caption_run.font.size = Pt(9)
    caption_run.font.italic = True
    caption_run.font.color.rgb = RGBColor(89, 97, 110)


def main():
    if not TARGET.exists():
        raise FileNotFoundError(TARGET)
    if list(IMAGE_DIR.glob("*.png")) if IMAGE_DIR.exists() else False:
        raise FileExistsError(f"Refusing to overwrite existing manual images in {IMAGE_DIR}")

    document = Document(TARGET)
    if document.inline_shapes:
        raise ValueError("The manual already contains inline images; refusing to add duplicate figures.")

    placements = []
    for source_name, anchor_prefix, caption in FIGURES:
        source_path = SOURCE / source_name
        if not source_path.exists():
            raise FileNotFoundError(source_path)
        target_name = f"figure-{len(placements) + 1:02d}-{source_name.split('_', 1)[1].replace('_', '-')}"
        placements.append((source_path, anchor_prefix, caption, target_name))

    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    for source_path, anchor_prefix, caption, target_name in placements:
        shutil.copy2(source_path, IMAGE_DIR / target_name)

    # Resolve all anchors before inserting paragraphs so the source text remains easy to match.
    resolved = [
        (matching_paragraph(document, anchor_prefix), IMAGE_DIR / target_name, caption)
        for _, anchor_prefix, caption, target_name in placements
    ]
    for number, (anchor, image_path, caption) in enumerate(resolved, start=1):
        add_figure_after(anchor, image_path, number, caption)

    document.save(TARGET)
    print(f"Embedded {len(resolved)} figures in {TARGET}")
    print(f"Copied source images to {IMAGE_DIR}")


if __name__ == "__main__":
    main()
